#!/usr/bin/env python3
"""
El cerebro de Luna: conversa y ejecuta tareas usando la API gratuita de Groq
(modelos GPT-OSS, sin costo) con "function calling" (herramientas).

No requiere pagar nada. Solo necesitas una API key gratis de https://console.groq.com
(sección "API Keys"), puesta en la variable de entorno GROQ_API_KEY (o en un
archivo .env junto a este script).

Uso rápido:
    from cerebro import Cerebro
    luna = Cerebro()
    respuesta = luna.procesar("¿qué hora es?")
    print(respuesta)

Herramientas incluidas (todas gratis, sin clave adicional):
    - fecha y hora actual
    - clima de cualquier ciudad (Open-Meteo)
    - búsqueda en internet (DuckDuckGo)
    - crear un informe/documento en PDF, Word, Excel o texto plano
    - guardar/consultar/completar recordatorios y pendientes (memoria.py,
      persiste entre reinicios)
    - abrir una página web o una carpeta/app en Windows
    - poner un temporizador que avisa por voz al terminar
    - leer lo que hay copiado en el portapapeles
    - subir/bajar/silenciar el volumen del sistema (Windows)
"""

from __future__ import annotations

import json
import logging
import os
import subprocess
import sys
import threading
import time
import webbrowser
from datetime import datetime
from pathlib import Path
from typing import Any, Callable

from dotenv import load_dotenv

load_dotenv(Path(__file__).resolve().parent / ".env")

log = logging.getLogger("cerebro")

# Groq va cambiando/retirando nombres de modelo con el tiempo. Si en algún momento
# vuelve a fallar con "model_not_found", corre esto en una terminal (con tu clave)
# para ver los que están disponibles ahora y soportan "tools":
#   curl https://api.groq.com/openai/v1/models -H "Authorization: Bearer TU_CLAVE"
# y pon el nombre que quieras en GROQ_MODEL dentro del .env.
GROQ_MODEL = (os.environ.get("GROQ_MODEL") or "openai/gpt-oss-120b").strip()
GROQ_MODEL_RESPALDO = (
    os.environ.get("GROQ_MODEL_RESPALDO") or "openai/gpt-oss-20b"
).strip()
MAX_TURNOS_HERRAMIENTA = 6  # evita bucles infinitos de llamadas a herramientas
MAX_MENSAJES_HISTORIAL = 20  # recorta el historial para no gastar tokens de más

NOMBRE_USUARIA = (os.environ.get("JARVIS_USUARIA") or "Daniela").strip()
# Nombre con el que la IA se identifica al hablar. Cámbialo en el .env
# (JARVIS_NOMBRE_ASISTENTE=OtroNombre) si algún día quieres otro distinto a "Luna".
NOMBRE_ASISTENTE = (os.environ.get("JARVIS_NOMBRE_ASISTENTE") or "Luna").strip()

PROMPT_SISTEMA = f"""Eres {NOMBRE_ASISTENTE.upper()}, un asistente de inteligencia artificial
profesional y capaz, al estilo del asistente de Iron Man. Hablas en español, de forma
natural, directa y un poco elegante, sin relleno ni disculpas innecesarias. Te diriges
a tu usuaria por su nombre, {NOMBRE_USUARIA}, cuando tenga sentido, sin abusar de
repetirlo.

Reglas:
- Si la pregunta se responde con una herramienta (fecha, hora, clima, búsqueda en
  internet, crear un documento/informe, abrir algo, temporizador, portapapeles,
  volumen, recordatorios), ÚSALA en vez de inventar la respuesta.
- Tienes memoria persistente de pendientes (recordatorios, llamadas por hacer,
  tareas) que sobrevive entre reinicios. Cuando la usuaria te pida que recuerdes
  algo, guárdalo con crear_recordatorio. Cuando pregunte qué tiene pendiente,
  usa listar_pendientes. Cuando diga que ya hizo algo, márcalo con
  completar_pendiente.
- Si no hace falta ninguna herramienta (charla, opinión, explicación, cálculo,
  redacción, código, consejo, etc.), respóndelo tú directamente con tu propio
  conocimiento, con la misma calidad y profundidad que darías en cualquier tema,
  igual que lo haría un asistente profesional serio, no una respuesta básica o
  genérica.
- Sé conciso cuando la respuesta se va a leer en voz alta (unas pocas frases), y más
  detallado cuando la pregunta claramente pide detalle o la orden llegó escrita.
- Nunca digas que "no tienes cerebro" ni que te falta información básica: si algo no
  lo sabes con certeza, dilo claramente y ofrece cómo averiguarlo (por ejemplo, con
  la herramienta de búsqueda en internet).
"""


# --------------------------------------------------------------------------- #
# Herramientas (funciones reales que Luna puede ejecutar)
# --------------------------------------------------------------------------- #

def _carpeta_documentos_jarvis() -> Path:
    base = Path(os.environ.get("USERPROFILE") or Path.home())
    carpeta = base / "Documents" / "Luna"
    carpeta.mkdir(parents=True, exist_ok=True)
    return carpeta


def herramienta_fecha_hora(**_: Any) -> dict:
    ahora = datetime.now()
    dias = [
        "lunes", "martes", "miércoles", "jueves", "viernes", "sábado", "domingo",
    ]
    meses = [
        "enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto",
        "septiembre", "octubre", "noviembre", "diciembre",
    ]
    texto = (
        f"Hoy es {dias[ahora.weekday()]} {ahora.day} de {meses[ahora.month - 1]} de "
        f"{ahora.year}, son las {ahora.strftime('%I:%M %p')}."
    )
    return {"texto": texto, "iso": ahora.isoformat()}


def herramienta_clima(ciudad: str, **_: Any) -> dict:
    """Clima actual con Open-Meteo: 100% gratis, sin API key."""
    import requests

    geo = requests.get(
        "https://geocoding-api.open-meteo.com/v1/search",
        params={"name": ciudad, "count": 1, "language": "es"},
        timeout=10,
    ).json()
    resultados = geo.get("results") or []
    if not resultados:
        return {"error": f"No encontré la ciudad «{ciudad}»."}
    lugar = resultados[0]
    lat, lon = lugar["latitude"], lugar["longitude"]

    clima = requests.get(
        "https://api.open-meteo.com/v1/forecast",
        params={
            "latitude": lat,
            "longitude": lon,
            "current": "temperature_2m,relative_humidity_2m,weather_code,wind_speed_10m",
            "timezone": "auto",
        },
        timeout=10,
    ).json()
    actual = clima.get("current") or {}
    codigos = {
        0: "cielo despejado", 1: "mayormente despejado", 2: "parcialmente nublado",
        3: "nublado", 45: "niebla", 48: "niebla helada", 51: "llovizna ligera",
        53: "llovizna", 55: "llovizna intensa", 61: "lluvia ligera", 63: "lluvia",
        65: "lluvia fuerte", 71: "nieve ligera", 73: "nieve", 75: "nieve fuerte",
        80: "chubascos ligeros", 81: "chubascos", 82: "chubascos fuertes",
        95: "tormenta", 96: "tormenta con granizo", 99: "tormenta fuerte con granizo",
    }
    descripcion = codigos.get(actual.get("weather_code"), "condiciones variables")
    nombre_lugar = f"{lugar.get('name')}, {lugar.get('country', '')}".strip(", ")
    texto = (
        f"En {nombre_lugar} hay {descripcion}, la temperatura es de "
        f"{actual.get('temperature_2m')}°C, humedad {actual.get('relative_humidity_2m')}% "
        f"y viento de {actual.get('wind_speed_10m')} km/h."
    )
    return {"texto": texto, "datos": actual, "lugar": nombre_lugar}


def herramienta_buscar_internet(consulta: str, **_: Any) -> dict:
    """Búsqueda web gratis con DuckDuckGo, sin clave."""
    try:
        from ddgs import DDGS
    except ImportError:
        return {"error": "Falta instalar ddgs (pip install -r requirements.txt)."}

    resultados = []
    try:
        with DDGS() as ddgs:
            for r in ddgs.text(consulta, region="es-es", max_results=5):
                resultados.append(
                    {
                        "titulo": r.get("title"),
                        "resumen": r.get("body"),
                        "url": r.get("href"),
                    }
                )
    except Exception as e:  # noqa: BLE001 - queremos degradar sin tumbar el turno
        return {"error": f"La búsqueda falló: {e}"}

    if not resultados:
        return {"error": "No encontré resultados para esa búsqueda."}
    return {"resultados": resultados}


def _texto_pdf_seguro(texto: str) -> str:
    """La fuente base del PDF solo soporta latin-1: quita emojis y símbolos raros en
    vez de tumbar la generación del PDF."""
    return texto.encode("latin-1", "ignore").decode("latin-1")


def _slug(texto: str) -> str:
    limpio = "".join(c if c.isalnum() or c in " _-" else "" for c in texto)
    return limpio.strip().replace(" ", "_")[:60]


def _generar_pdf(ruta: Path, titulo: str, contenido: str) -> None:
    from fpdf import FPDF
    from fpdf.enums import XPos, YPos

    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 16)
    pdf.multi_cell(0, 10, _texto_pdf_seguro(titulo), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.ln(4)
    pdf.set_font("Helvetica", "", 12)
    for linea in contenido.splitlines() or [contenido]:
        pdf.multi_cell(0, 8, _texto_pdf_seguro(linea) or " ", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.output(str(ruta))


def _generar_docx(ruta: Path, titulo: str, contenido: str) -> None:
    from docx import Document

    doc = Document()
    doc.add_heading(titulo, level=1)
    for linea in contenido.splitlines() or [contenido]:
        doc.add_paragraph(linea)
    doc.save(str(ruta))


def _generar_xlsx(ruta: Path, titulo: str, contenido: str) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = (titulo[:31] or "Hoja1")  # Excel limita el nombre de hoja a 31 caracteres
    ws.append([titulo])
    ws.append([])
    # Cada línea del contenido es una fila; si trae comas, cada parte va en una
    # columna distinta (así "nombre, cantidad, precio" queda en 3 celdas).
    for linea in contenido.splitlines() or [contenido]:
        celdas = [c.strip() for c in linea.split(",")] if "," in linea else [linea]
        ws.append(celdas)
    wb.save(str(ruta))


def _generar_txt(ruta: Path, titulo: str, contenido: str) -> None:
    ruta.write_text(f"{titulo}\n{'=' * len(titulo)}\n\n{contenido}\n", encoding="utf-8")


_GENERADORES_DOCUMENTO = {
    "pdf": (_generar_pdf, "pdf"),
    "docx": (_generar_docx, "docx"),
    "word": (_generar_docx, "docx"),
    "xlsx": (_generar_xlsx, "xlsx"),
    "excel": (_generar_xlsx, "xlsx"),
    "txt": (_generar_txt, "txt"),
    "texto": (_generar_txt, "txt"),
}


def herramienta_crear_documento(
    titulo: str, contenido: str, formato: str = "pdf", **_: Any
) -> dict:
    """Genera un documento/informe en el formato que pida la usuaria (pdf, Word,
    Excel o texto plano) y lo guarda en Documentos/Luna."""
    formato_normalizado = (formato or "pdf").strip().lower()
    entrada = _GENERADORES_DOCUMENTO.get(formato_normalizado)
    if entrada is None:
        formatos = ", ".join(sorted({v[1] for v in _GENERADORES_DOCUMENTO.values()}))
        return {"error": f"Formato «{formato}» no soportado. Usa uno de: {formatos}."}
    generador, extension = entrada

    carpeta = _carpeta_documentos_jarvis()
    marca_tiempo = datetime.now().strftime("%Y-%m-%d_%H%M")
    nombre_archivo = f"{_slug(titulo) or 'documento'}_{marca_tiempo}.{extension}"
    ruta = carpeta / nombre_archivo
    try:
        generador(ruta, titulo, contenido)
    except ImportError as e:
        return {"error": f"Falta instalar una librería para generar .{extension}: {e}"}
    return {"texto": f"Listo, guardé el documento en {ruta}", "ruta": str(ruta)}


def herramienta_abrir(destino: str, **_: Any) -> dict:
    """Abre una URL, o en Windows una carpeta/archivo/app, con el programa por defecto."""
    destino = destino.strip()
    if not destino:
        return {"error": "No me diste qué abrir."}
    if destino.startswith(("http://", "https://")):
        webbrowser.open(destino)
        return {"texto": f"Abriendo {destino}"}
    if sys.platform == "win32":
        try:
            os.startfile(destino)  # type: ignore[attr-defined]
            return {"texto": f"Abriendo {destino}"}
        except OSError as e:
            return {"error": f"No pude abrir «{destino}»: {e}"}
    # Fallback no-Windows: intenta como sitio web con https:// al frente.
    webbrowser.open(f"https://{destino}")
    return {"texto": f"Abriendo {destino}"}


def herramienta_temporizador(segundos: float, mensaje: str, hablar_fn: Callable[[str], None], **_: Any) -> dict:
    segundos = max(1.0, float(segundos))
    aviso = mensaje.strip() or "Se acabó el tiempo."

    def _avisar() -> None:
        time.sleep(segundos)
        hablar_fn(aviso)

    threading.Thread(target=_avisar, daemon=True).start()
    if segundos < 60:
        return {"texto": f"Temporizador puesto para {segundos:.0f} segundos."}
    return {"texto": f"Temporizador puesto para {segundos / 60:.1f} minutos."}


def herramienta_portapapeles(**_: Any) -> dict:
    try:
        import pyperclip
    except ImportError:
        return {"error": "Falta instalar pyperclip (pip install -r requirements.txt)."}
    try:
        contenido = pyperclip.paste()
    except Exception as e:  # noqa: BLE001
        return {"error": f"No pude leer el portapapeles: {e}"}
    if not contenido or not contenido.strip():
        return {"texto": "El portapapeles está vacío."}
    return {"texto": contenido.strip()[:4000]}


def herramienta_volumen(accion: str, nivel: int | None = None, **_: Any) -> dict:
    """Control de volumen del sistema (solo Windows, vía pycaw)."""
    if sys.platform != "win32":
        return {"error": "El control de volumen solo está disponible en Windows."}
    try:
        from ctypes import cast, POINTER
        from comtypes import CLSCTX_ALL
        from pycaw.pycaw import AudioUtilities, IAudioEndpointVolume
    except ImportError:
        return {"error": "Falta instalar pycaw y comtypes (pip install -r requirements.txt)."}

    dispositivos = AudioUtilities.GetSpeakers()
    interfaz = dispositivos.Activate(IAudioEndpointVolume._iid_, CLSCTX_ALL, None)
    volumen = cast(interfaz, POINTER(IAudioEndpointVolume))

    accion = accion.lower().strip()
    if accion == "silenciar":
        volumen.SetMute(1, None)
        return {"texto": "Volumen silenciado."}
    if accion == "activar":
        volumen.SetMute(0, None)
        return {"texto": "Sonido activado."}
    if accion == "fijar" and nivel is not None:
        nivel = max(0, min(100, int(nivel)))
        volumen.SetMasterVolumeLevelScalar(nivel / 100, None)
        return {"texto": f"Volumen en {nivel}%."}
    actual = volumen.GetMasterVolumeLevelScalar()
    paso = 0.1
    if accion == "subir":
        volumen.SetMasterVolumeLevelScalar(min(1.0, actual + paso), None)
    elif accion == "bajar":
        volumen.SetMasterVolumeLevelScalar(max(0.0, actual - paso), None)
    else:
        return {"error": f"Acción de volumen desconocida: {accion}"}
    nuevo = round(volumen.GetMasterVolumeLevelScalar() * 100)
    return {"texto": f"Volumen en {nuevo}%."}


def herramienta_crear_recordatorio(
    texto: str, fecha_hora: str | None = None, tipo: str = "recordatorio", **_: Any
) -> dict:
    """Guarda un pendiente (recordatorio, llamada por hacer, tarea) que persiste
    entre reinicios. Si trae fecha_hora, además se avisa por voz a esa hora."""
    import memoria

    vence_iso = None
    if fecha_hora and fecha_hora.strip():
        try:
            vence_iso = datetime.fromisoformat(fecha_hora.strip()).isoformat(timespec="seconds")
        except ValueError:
            return {
                "error": (
                    f"No entendí la fecha/hora «{fecha_hora}». Usa formato ISO, ej: "
                    "2026-09-15T15:00:00 (usa la herramienta fecha_hora si necesitas saber hoy)."
                )
            }
    item = memoria.agregar_pendiente(texto, vence=vence_iso, tipo=tipo)
    cuando = f" para el {vence_iso}" if vence_iso else ""
    return {"texto": f"Guardado: «{item['texto']}»{cuando}.", "id": item["id"]}


def herramienta_listar_pendientes(**_: Any) -> dict:
    import memoria

    return {"texto": memoria.resumen_pendientes_hoy()}


def herramienta_completar_pendiente(referencia: str, **_: Any) -> dict:
    """referencia: el id del pendiente, o parte del texto (ej. 'contador')."""
    import memoria

    item = memoria.completar_pendiente(referencia)
    if item is None:
        return {"error": f"No encontré ningún pendiente que coincida con «{referencia}»."}
    return {"texto": f"Marcado como hecho: «{item['texto']}»."}


# --------------------------------------------------------------------------- #
# Definición de herramientas en el formato "function calling" (OpenAI-compatible,
# que es el que usa la API de Groq).
# --------------------------------------------------------------------------- #

_HERRAMIENTAS_ESQUEMA = [
    {
        "type": "function",
        "function": {
            "name": "fecha_hora",
            "description": "Devuelve la fecha y hora actual exacta.",
            "parameters": {"type": "object", "properties": {}},
        },
    },
    {
        "type": "function",
        "function": {
            "name": "clima",
            "description": "Devuelve el clima actual de una ciudad.",
            "parameters": {
                "type": "object",
                "properties": {"ciudad": {"type": "string", "description": "Nombre de la ciudad"}},
                "required": ["ciudad"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "buscar_internet",
            "description": "Busca información actualizada en internet (noticias, precios, datos que cambian con el tiempo, o cualquier cosa que no sepas con certeza).",
            "parameters": {
                "type": "object",
                "properties": {"consulta": {"type": "string", "description": "Qué buscar"}},
                "required": ["consulta"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "crear_documento",
            "description": (
                "Crea un informe/documento con un título y contenido, y lo guarda en "
                "Documentos/Luna. Soporta varios formatos: usa el que pida la usuaria."
            ),
            "parameters": {
                "type": "object",
                "properties": {
                    "titulo": {"type": "string"},
                    "contenido": {
                        "type": "string",
                        "description": (
                            "Texto completo del documento, puede tener varias líneas. Para "
                            "xlsx, cada línea es una fila y las comas separan columnas."
                        ),
                    },
                    "formato": {
                        "type": "string",
                        "enum": ["pdf", "docx", "xlsx", "txt"],
                        "description": "Formato del archivo. Por defecto pdf si no se especifica.",
                    },
                },
                "required": ["titulo", "contenido"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "abrir",
            "description": "Abre una página web (con http/https) o, en Windows, una app/carpeta/archivo por su nombre o ruta.",
            "parameters": {
                "type": "object",
                "properties": {"destino": {"type": "string"}},
                "required": ["destino"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "temporizador",
            "description": "Pone un temporizador que avisa por voz al cumplirse el tiempo.",
            "parameters": {
                "type": "object",
                "properties": {
                    "segundos": {"type": "number", "description": "Duración en segundos"},
                    "mensaje": {"type": "string", "description": "Qué decir cuando termine"},
                },
                "required": ["segundos", "mensaje"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "portapapeles",
            "description": "Lee el texto que el usuario tiene copiado en el portapapeles.",
            "parameters": {"type": "object", "properties": {}},
        },
    },
    {
        "type": "function",
        "function": {
            "name": "volumen",
            "description": "Sube, baja, silencia, activa o fija el volumen del sistema (Windows).",
            "parameters": {
                "type": "object",
                "properties": {
                    "accion": {
                        "type": "string",
                        "enum": ["subir", "bajar", "silenciar", "activar", "fijar"],
                    },
                    "nivel": {"type": "integer", "description": "0-100, solo si accion es 'fijar'"},
                },
                "required": ["accion"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "crear_recordatorio",
            "description": (
                "Guarda un pendiente (recordatorio, algo por hacer, una llamada que hay que "
                "hacer, una tarea) que se recuerda incluso después de reiniciar. Si el usuario "
                "da una hora/fecha concreta, avisa por voz automáticamente cuando llegue."
            ),
            "parameters": {
                "type": "object",
                "properties": {
                    "texto": {"type": "string", "description": "Qué hay que recordar/hacer"},
                    "fecha_hora": {
                        "type": "string",
                        "description": (
                            "Opcional. Formato ISO 8601, ej: '2026-09-15T15:00:00'. Usa la "
                            "herramienta fecha_hora primero si necesitas saber la fecha de hoy "
                            "para calcular 'mañana', 'el viernes', etc. Si no se da, el "
                            "pendiente queda sin hora fija (solo aparece en la lista)."
                        ),
                    },
                    "tipo": {
                        "type": "string",
                        "enum": ["recordatorio", "llamada", "tarea"],
                        "description": "Clasificación opcional del pendiente.",
                    },
                },
                "required": ["texto"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "listar_pendientes",
            "description": "Devuelve todos los pendientes/recordatorios que faltan por completar.",
            "parameters": {"type": "object", "properties": {}},
        },
    },
    {
        "type": "function",
        "function": {
            "name": "completar_pendiente",
            "description": "Marca un pendiente como hecho/completado, por su id o por parte de su texto.",
            "parameters": {
                "type": "object",
                "properties": {
                    "referencia": {
                        "type": "string",
                        "description": "El id del pendiente, o una parte de su texto (ej. 'contador').",
                    }
                },
                "required": ["referencia"],
            },
        },
    },
]


class Cerebro:
    """Mantiene el historial de conversación y resuelve cada mensaje con Groq."""

    def __init__(self, hablar_fn: Callable[[str], None] | None = None) -> None:
        self._hablar_fn = hablar_fn or (lambda texto: log.info("Luna diría: %s", texto))
        self._historial: list[dict] = []
        self._lock = threading.Lock()
        self._cliente = None  # se crea de forma perezosa (lazy) al primer uso

        self._despachador: dict[str, Callable[..., dict]] = {
            "fecha_hora": herramienta_fecha_hora,
            "clima": herramienta_clima,
            "buscar_internet": herramienta_buscar_internet,
            "crear_documento": herramienta_crear_documento,
            "abrir": herramienta_abrir,
            "temporizador": lambda **kw: herramienta_temporizador(hablar_fn=self._hablar_fn, **kw),
            "portapapeles": herramienta_portapapeles,
            "volumen": herramienta_volumen,
            "crear_recordatorio": herramienta_crear_recordatorio,
            "listar_pendientes": herramienta_listar_pendientes,
            "completar_pendiente": herramienta_completar_pendiente,
        }

    # -- infraestructura ---------------------------------------------------- #

    def _obtener_cliente(self):
        if self._cliente is not None:
            return self._cliente
        clave = (os.environ.get("GROQ_API_KEY") or "").strip()
        if not clave:
            raise RuntimeError(
                "Falta GROQ_API_KEY. Crea una clave gratis en https://console.groq.com/keys "
                "y ponla en el archivo .env (GROQ_API_KEY=tu_clave)."
            )
        from groq import Groq

        # Por defecto el cliente espera hasta 10 minutos antes de rendirse; con un
        # límite corto, si algo se traba pasamos rápido al modelo de respaldo (o
        # avisamos el error) en vez de dejar el dashboard pegado en "pensando".
        self._cliente = Groq(api_key=clave, timeout=25.0, max_retries=1)
        return self._cliente

    def _recortar_historial(self) -> None:
        if len(self._historial) > MAX_MENSAJES_HISTORIAL:
            self._historial = self._historial[-MAX_MENSAJES_HISTORIAL:]

    def _ejecutar_herramienta(self, nombre: str, argumentos: dict) -> dict:
        funcion = self._despachador.get(nombre)
        if funcion is None:
            return {"error": f"Herramienta desconocida: {nombre}"}
        try:
            return funcion(**argumentos)
        except Exception as e:  # noqa: BLE001 - nunca tumbar el turno por una herramienta
            log.exception("Fallo ejecutando herramienta %s", nombre)
            return {"error": f"La herramienta {nombre} falló: {e}"}

    def _completar(self, cliente, modelo: str, mensajes: list[dict]):
        return cliente.chat.completions.create(
            model=modelo,
            messages=mensajes,
            tools=_HERRAMIENTAS_ESQUEMA,
            tool_choice="auto",
            temperature=0.5,
            max_tokens=1024,
        )

    # -- API pública ---------------------------------------------------------- #

    def procesar(self, mensaje_usuario: str) -> str:
        """Procesa un mensaje del usuario (voz o texto) y devuelve la respuesta de Luna."""
        mensaje_usuario = (mensaje_usuario or "").strip()
        if not mensaje_usuario:
            return "No escuché ninguna orden."

        with self._lock:
            try:
                cliente = self._obtener_cliente()
            except RuntimeError as e:
                return str(e)

            self._historial.append({"role": "user", "content": mensaje_usuario})
            mensajes = [{"role": "system", "content": PROMPT_SISTEMA}, *self._historial]

            modelos_a_probar = [GROQ_MODEL, GROQ_MODEL_RESPALDO]
            respuesta_final = None
            ultimo_error: Exception | None = None

            for modelo in modelos_a_probar:
                try:
                    respuesta_final = self._resolver_con_modelo(cliente, modelo, mensajes)
                    break
                except Exception as e:  # noqa: BLE001 - probamos el modelo de respaldo
                    ultimo_error = e
                    log.warning("Modelo %s falló (%s), probando respaldo...", modelo, e)
                    continue

            if respuesta_final is None:
                self._historial.pop()  # no dejamos un turno de usuario sin respuesta guardada
                return (
                    "No pude conectarme con el cerebro ahora mismo "
                    f"({ultimo_error}). Intenta de nuevo en un momento."
                )

            self._historial.append({"role": "assistant", "content": respuesta_final})
            self._recortar_historial()
            return respuesta_final

    def _resolver_con_modelo(self, cliente, modelo: str, mensajes: list[dict]) -> str:
        mensajes = list(mensajes)
        for _ in range(MAX_TURNOS_HERRAMIENTA):
            respuesta = self._completar(cliente, modelo, mensajes)
            eleccion = respuesta.choices[0].message
            llamadas = getattr(eleccion, "tool_calls", None)

            if not llamadas:
                return (eleccion.content or "").strip() or "Listo."

            mensajes.append(
                {
                    "role": "assistant",
                    "content": eleccion.content or "",
                    "tool_calls": [
                        {
                            "id": llamada.id,
                            "type": "function",
                            "function": {
                                "name": llamada.function.name,
                                "arguments": llamada.function.arguments,
                            },
                        }
                        for llamada in llamadas
                    ],
                }
            )
            for llamada in llamadas:
                try:
                    argumentos = json.loads(llamada.function.arguments or "{}")
                except json.JSONDecodeError:
                    argumentos = {}
                resultado = self._ejecutar_herramienta(llamada.function.name, argumentos)
                mensajes.append(
                    {
                        "role": "tool",
                        "tool_call_id": llamada.id,
                        "content": json.dumps(resultado, ensure_ascii=False),
                    }
                )

        return "Se me complicó resolver eso con tantos pasos; intenta pedírmelo de forma más simple."

    def historial_legible(self) -> list[dict]:
        with self._lock:
            return list(self._historial)


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO)
    luna = Cerebro()
    print("Cerebro de Luna (Groq) — escribe 'salir' para terminar.")
    while True:
        try:
            texto = input("Tú: ").strip()
        except (EOFError, KeyboardInterrupt):
            break
        if texto.lower() in {"salir", "exit", "quit"}:
            break
        print("Luna:", luna.procesar(texto))
